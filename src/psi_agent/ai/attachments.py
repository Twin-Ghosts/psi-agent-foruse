"""把用户消息里的 ``[RECV:/path]`` 附件标记展开成模型能消费的内容。

链路背景: Gateway 收到上传文件后存盘、以 ``[RECV:/绝对路径]`` 文本标记随消息发来,
但此前无人「读文件注入内容」, 模型只看到一串路径。这里补上这一步:

- **模型能理解该文件** → 原样给 (图片走多模态 ``image_url`` 块, 交模型自己看) 。
- **模型理解不了** → 转成文字 (pdf/docx/xlsx/文本抽取正文; 图片退化为占位说明) 。

判断「能否理解」用模型名启发式: any-llm 无离线能力查询, 故按已知支持视觉的模型族匹配,
未知一律当纯文本模型 (宁可转文字, 也不塞模型看不懂的东西) 。
"""

from __future__ import annotations

import base64
import mimetypes
import os
import re
from typing import Any

from loguru import logger

RECV_RE = re.compile(r"\[\s*RECV\s*:\s*(.+?)\s*\]", re.IGNORECASE)

# 单个附件注入上限, 防止超大文件把上下文撑爆 (字符数, 约等于 token 的几倍)
_MAX_TEXT_CHARS = 60_000
# 图片走多模态时的字节上限 (约 8MB, base64 后更大)
_MAX_IMAGE_BYTES = 8 * 1024 * 1024

_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}

# 已知支持图像输入的模型族 (子串匹配, 小写) 。未命中一律按纯文本处理。
_VISION_HINTS = (
    "gpt-4o", "gpt-4.1", "gpt-4-vision", "gpt-4-turbo", "o1", "o3", "o4",
    "claude-3", "claude-4", "claude-opus", "claude-sonnet", "claude-haiku",
    "gemini", "qwen-vl", "qwen2-vl", "qwen2.5-vl", "llava", "pixtral",
    "grok-vision", "grok-2-vision", "internvl", "glm-4v", "step-1v",
)


def model_supports_vision(provider: str, model: str) -> bool:
    """按模型名启发式判断是否支持图像输入。未知按不支持。"""
    m = (model or "").lower()
    return any(h in m for h in _VISION_HINTS)


def _extract_text(path: str) -> str:
    """按扩展名抽取文件正文文本。失败/不支持时返回空串 (由调用方给占位) 。"""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            import pymupdf  # type: ignore  # noqa: PLC0415

            doc = pymupdf.open(path)
            try:
                return "\n".join(page.get_text() for page in doc)
            finally:
                doc.close()
        if ext == ".docx":
            import docx  # type: ignore  # noqa: PLC0415

            return "\n".join(p.text for p in docx.Document(path).paragraphs)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl  # type: ignore  # noqa: PLC0415

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows: list[str] = []
            for ws in wb.worksheets:
                rows.append(f"# 工作表: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            wb.close()
            return "\n".join(rows)
        # 其余按纯文本读 (txt/md/csv/json/代码等)
        with open(path, "rb") as f:
            raw = f.read(_MAX_TEXT_CHARS * 4)
        return raw.decode("utf-8", errors="replace")
    except Exception as e:
        logger.warning(f"附件抽取失败 {path!r}: {e!r}")
        return ""


def _read_image_block(path: str) -> dict[str, Any] | None:
    """把图片读成 OpenAI 多模态 image_url 块 (data URI) 。过大/失败返回 None。"""
    try:
        size = os.path.getsize(path)
        if size > _MAX_IMAGE_BYTES:
            logger.warning(f"图片过大跳过多模态 {path!r} ({size} bytes)")
            return None
        mime = mimetypes.guess_type(path)[0] or "image/png"
        with open(path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode()
        return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}
    except Exception as e:
        logger.warning(f"图片读取失败 {path!r}: {e!r}")
        return None


def _text_note(path: str) -> str:
    """把一个附件转成注入正文的文字块。"""
    name = os.path.basename(path)
    if not os.path.exists(path):
        return f"\n\n[附件 {name}: 文件不存在或已被移动, 无法读取]\n"
    ext = os.path.splitext(path)[1].lower()
    if ext in _IMAGE_EXT:
        return f"\n\n[图片附件 {name}: 当前模型不支持图像输入, 无法查看图片内容]\n"
    text = _extract_text(path)
    if not text.strip():
        return f"\n\n[附件 {name}: 无法提取文本内容]\n"
    if len(text) > _MAX_TEXT_CHARS:
        text = text[:_MAX_TEXT_CHARS] + f"\n… (内容过长, 已截断, 共 {len(text)} 字符) "
    return f"\n\n[附件 {name} 内容开始]\n{text}\n[附件 {name} 内容结束]\n"


def _expand_one(content: Any, vision: bool) -> Any:
    """展开单条消息 content 里的 [RECV:] 标记。

    - 无标记: 原样返回。
    - 有图片且模型支持视觉: content 变多模态数组 (文字块 + image_url 块) 。
    - 其余: 标记就地替换为抽取的文字, content 仍是字符串。
    """
    if not isinstance(content, str) or "[RECV" not in content.upper():
        return content
    paths = [m.group(1).strip() for m in RECV_RE.finditer(content)]
    if not paths:
        return content

    image_blocks: list[dict[str, Any]] = []
    if vision:
        for p in paths:
            if os.path.splitext(p)[1].lower() in _IMAGE_EXT and os.path.exists(p):
                blk = _read_image_block(p)
                if blk is not None:
                    image_blocks.append(blk)

    def _sub(match: re.Match[str]) -> str:
        p = match.group(1).strip()
        name = os.path.basename(p)
        if vision and os.path.splitext(p)[1].lower() in _IMAGE_EXT and os.path.exists(p):
            return f"\n\n[图片附件 {name} (见下方图片) ]\n"
        return _text_note(p)

    text_content = RECV_RE.sub(_sub, content)
    if image_blocks:
        return [{"type": "text", "text": text_content}, *image_blocks]
    return text_content


_FILE_PROTOCOL_MARK = "文件传输协议: "

# 说明文字里**刻意不写出完整的标记字面量**: SEND_RE 是 `\[\s*SEND\s*:\s*(.+?)\s*\]`,
# 见到就当成一次真交付。模型复述协议说明是很常见的行为 ("你可以用 [SEND:/路径]
# 这样的格式…"), 一旦复述就会触发一次指向不存在路径的交付。所以这里把标记拆开
# 描述, 示例也用真实感的路径而非 "/绝对/路径" 这种占位串。
_FILE_PROTOCOL = (
    _FILE_PROTOCOL_MARK
    + "当你需要把一个文件交付给用户 (生成的文档、表格、图片、报告等), "
    "在回复正文里写一行标记: 左方括号 + SEND: + 该文件的绝对路径 + 右方括号。"
    "例如文件在 D:/work/report.xlsx, 就写 SEND: 冒号后接这个路径并用方括号括起。"
    "系统会据此把文件交付给用户、在界面按会话累计为「交付物」。"
    "务必先用工具真实创建该文件再引用其路径; 一个回复可含多个这样的标记。"
    "注意: 只在真的要交付文件时才写这个标记, "
    "解释或举例说明这个格式时不要写出完整标记, 否则会触发一次无效交付。"
    "用户消息里出现的 RECV 标记是用户上传的附件 (内容已为你展开), 无需你处理。"
)


def inject_file_protocol(messages: list[Any]) -> list[Any]:
    """确保系统消息里含 [SEND:] 文件交付协议说明。

    默认 agent 的 system prompt 为空, 模型不知道怎么交付文件, 导致「交付物」永远
    不累计。这里幂等注入: 已有 system 消息则追加 (若尚未含协议), 否则插一条。

    去重只看 **system 消息里的协议标记**, 不能扫全部消息找 "[SEND:": 助手历史里
    只要交付过一次文件就带着这个串, 那样第二轮起会被误判成"讲过了"而跳过注入,
    模型反而在多轮会话中失去协议说明 —— 与本函数的目的正好相反。
    """
    if not messages:
        return messages
    for m in messages:
        if (isinstance(m, dict) and m.get("role") == "system"
                and isinstance(m.get("content"), str)
                and _FILE_PROTOCOL_MARK in m["content"]):
            return messages
    out = list(messages)
    first = out[0] if out else None
    if isinstance(first, dict) and first.get("role") == "system" and isinstance(first.get("content"), str):
        merged = (first["content"] + "\n\n" + _FILE_PROTOCOL) if first["content"].strip() else _FILE_PROTOCOL
        out[0] = {**first, "content": merged}
    else:
        out.insert(0, {"role": "system", "content": _FILE_PROTOCOL})
    return out


def expand_attachments(messages: list[Any], provider: str, model: str) -> list[Any]:
    """把 messages 里所有 user 消息的 [RECV:] 标记展开成模型可消费的内容。

    非破坏: 无标记的消息原样返回。历史里的附件同样展开 (模型应看得到) 。
    """
    if not messages:
        return messages
    vision = model_supports_vision(provider, model)
    out: list[Any] = []
    for msg in messages:
        if isinstance(msg, dict) and msg.get("role") == "user" and "content" in msg:
            expanded = _expand_one(msg["content"], vision)
            if expanded is not msg["content"]:
                msg = {**msg, "content": expanded}
        out.append(msg)
    return out
